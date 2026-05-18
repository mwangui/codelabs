from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "DOC"
CHART_DIR = DOC_DIR / "charts"
OUT_DOCX = DOC_DIR / "01_Google_Codelabs_學習地圖與完整索引.docx"
OUT_CSV = DOC_DIR / "01_Google_Codelabs_完整索引.csv"


CATEGORY_ORDER = {
    "AI / Gemini / Agent / Vertex AI": 1,
    "資料分析 / BigQuery / Looker": 2,
    "Cloud Run / Serverless / App Engine": 3,
    "Kubernetes / GKE / Containers / DevOps": 4,
    "網路 / 安全 / 身分 / 基礎架構": 5,
    "Firebase / Web / Flutter / Mobile": 6,
    "Android / Chrome / Web Platform": 7,
    "Workspace / Assistant / Dialogflow / Smart Home": 8,
    "ML / Vision / Document AI / Media / AR": 9,
    "Payments / Wallet / Commerce": 10,
    "Health / Maps / IoT / 其他": 11,
}

LEVEL_ORDER = {"初階": 1, "中階": 2, "進階": 3, "高階": 4, "未分類": 9}


@dataclass
class Record:
    index: int
    slug: str
    title: str
    url: str
    pdf_filename: str
    status: str
    step_count: int
    pdf_page_count: int
    updated_date: str
    duration_text: str
    duration_minutes: int
    description: str
    category: str
    level: str
    ai_subtopic: str


def clean_int(value: str | None) -> int:
    try:
        return int(value or 0)
    except ValueError:
        return 0


def parse_updated_date(time_text: str) -> str:
    match = re.search(r"Updated\s+(\d{4})年(\d{1,2})月(\d{1,2})日", time_text or "")
    if not match:
        return ""
    return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"


def parse_duration_from_text(title: str, text: str) -> tuple[str, int, str]:
    duration_text = ""
    description = ""
    if text.startswith(title) and "Updated" in text:
        after_title = text[len(title) :]
        before_updated, after_updated = after_title.split("Updated", 1)
        duration_text = before_updated.strip()
        description = re.sub(r"^\s*\d{4}年\d{1,2}月\d{1,2}日", "", after_updated).strip()
        if description.endswith("Start"):
            description = description[:-5].strip()

    minutes = 0
    hour_match = re.search(r"(\d+)\s*小時", duration_text)
    minute_match = re.search(r"(\d+)\s*分鐘", duration_text)
    if hour_match:
        minutes += int(hour_match.group(1)) * 60
    if minute_match:
        minutes += int(minute_match.group(1))
    return duration_text, minutes, description


def contains_any(text: str, words: Iterable[str]) -> bool:
    lower = text.lower()
    return any(word.lower() in lower for word in words)


def classify_category(title: str, slug: str, description: str) -> str:
    text = f"{title} {slug} {description}"
    if contains_any(
        text,
        [
            "gemini",
            "gen ai",
            "genai",
            "generative ai",
            "生成式",
            "agent",
            "代理",
            "adk",
            "vertex ai",
            "llm",
            "rag",
            "prompt",
            "mcp",
            "model armor",
            "agentspace",
            "genui",
        ],
    ):
        return "AI / Gemini / Agent / Vertex AI"
    if contains_any(text, ["bigquery", "looker", "dataflow", "data fusion", "dataproc", "pub/sub", "analytics", "資料分析", "資料集"]):
        return "資料分析 / BigQuery / Looker"
    if contains_any(text, ["cloud run", "cloud functions", "functions", "app engine", "serverless", "cloud buildpacks", "buildpacks"]):
        return "Cloud Run / Serverless / App Engine"
    if contains_any(text, ["gke", "kubernetes", "kubeflow", "container", "docker", "jib", "skaffold", "devops", "cloud build", "cloud deploy", "istio"]):
        return "Kubernetes / GKE / Containers / DevOps"
    if contains_any(
        text,
        [
            "security",
            "安全",
            "identity",
            "iam",
            "network",
            "網路",
            "cloud armor",
            "cloud dns",
            "private service connect",
            "psc",
            "kms",
            "vpc",
            "firewall",
            "tls",
            "load balanc",
        ],
    ):
        return "網路 / 安全 / 身分 / 基礎架構"
    if contains_any(text, ["firebase", "flutter", "mobile", "react native", "web app", "pwa", "lit", "angular", "web"]):
        return "Firebase / Web / Flutter / Mobile"
    if contains_any(text, ["android", "chrome", "web bluetooth", "webgpu", "web platform", "amp", "arcore"]):
        return "Android / Chrome / Web Platform"
    if contains_any(text, ["dialogflow", "assistant", "actions", "smart home", "workspace", "gmail", "calendar", "chatbot"]):
        return "Workspace / Assistant / Dialogflow / Smart Home"
    if contains_any(text, ["vision", "document ai", "ml", "machine learning", "automl", "tensorflow", "tensorflow.js", "lite", "ar", "media", "image"]):
        return "ML / Vision / Document AI / Media / AR"
    if contains_any(text, ["pay", "wallet", "commerce", "admob", "ads"]):
        return "Payments / Wallet / Commerce"
    return "Health / Maps / IoT / 其他"


def classify_ai_subtopic(title: str, slug: str, description: str, category: str) -> str:
    if category != "AI / Gemini / Agent / Vertex AI":
        return ""
    text = f"{title} {slug} {description}"
    if contains_any(text, ["agent gateway", "agentspace", "agent platform", "enterprise agent"]):
        return "Enterprise Agent Platform / Gateway"
    if contains_any(text, ["multi-agent", "多代理", "adk", "agents cli", "agent development kit", "agentic"]):
        return "ADK / 多代理 / Agentic Apps"
    if contains_any(text, ["mcp", "workspace"]):
        return "MCP / Workspace 代理"
    if contains_any(text, ["evaluation", "評估", "observability", "觀測"]):
        return "AI 評估 / 觀測"
    if contains_any(text, ["security", "安全", "model armor", "governance", "治理"]):
        return "AI 安全 / 治理"
    if contains_any(text, ["rag", "retrieval", "檢索", "search", "ground"]):
        return "RAG / grounding / 搜尋"
    if contains_any(text, ["bigquery", "data agent", "資料代理"]):
        return "BigQuery AI / 資料代理"
    if contains_any(text, ["vertex ai"]):
        return "Vertex AI / Gemini API"
    if contains_any(text, ["gemini"]):
        return "Gemini App 開發"
    return "其他 AI / GenAI"


def classify_level(title: str, slug: str, description: str, duration_minutes: int, steps: int, pages: int) -> str:
    text = f"{title} {slug} {description}"
    if contains_any(
        text,
        [
            "enterprise",
            "production-ready",
            "production ready",
            "gateway",
            "multi-agent",
            "多代理",
            "security",
            "安全",
            "observability",
            "evaluation",
            "評估",
            "networking",
            "private service connect",
            "istio",
            "kubeflow",
            "migration",
            "遷移",
            "advanced",
            "進階",
            "multi-cloud",
        ],
    ) or duration_minutes >= 100 or steps >= 16 or pages >= 34:
        return "高階"
    if contains_any(
        text,
        [
            "deploy",
            "部署",
            "build",
            "建構",
            "integrat",
            "整合",
            "pipeline",
            "rag",
            "vertex ai",
            "bigquery",
            "cloud run",
            "firebase",
            "gke",
            "grpc",
        ],
    ) or duration_minutes >= 60 or steps >= 10 or pages >= 22:
        return "進階"
    if contains_any(text, ["overview", "總覽", "intro", "簡介", "getting started", "開始使用", "quickstart", "first", "基礎", "入門", "hello"]) or duration_minutes <= 35:
        return "初階"
    return "中階"


def load_records() -> list[Record]:
    report_rows = list(csv.DictReader((ROOT / "output" / "final_report.csv").open(encoding="utf-8")))
    metadata_rows = json.loads((ROOT / "output" / "catalog_metadata.json").read_text(encoding="utf-8"))
    metadata_by_url = {row["url"]: row for row in metadata_rows}

    records: list[Record] = []
    for row in report_rows:
        meta = metadata_by_url.get(row["url"], {})
        title = row["title"] or meta.get("title", "")
        updated_date = parse_updated_date(str(meta.get("time_text", "")))
        duration_text, duration_minutes, description = parse_duration_from_text(title, str(meta.get("text", "")))
        category = classify_category(title, row["slug"], description)
        ai_subtopic = classify_ai_subtopic(title, row["slug"], description, category)
        steps = clean_int(row.get("step_count"))
        pages = clean_int(row.get("pdf_page_count"))
        level = classify_level(title, row["slug"], description, duration_minutes, steps, pages)
        records.append(
            Record(
                index=clean_int(row["index"]),
                slug=row["slug"],
                title=title,
                url=row["url"],
                pdf_filename=row["pdf_filename"],
                status=row["status"],
                step_count=steps,
                pdf_page_count=pages,
                updated_date=updated_date,
                duration_text=duration_text,
                duration_minutes=duration_minutes,
                description=description,
                category=category,
                level=level,
                ai_subtopic=ai_subtopic,
            )
        )
    return records


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 8.0, color: str = "202124") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    return paragraph


def add_body_paragraph(doc: Document, text: str, *, bold_label: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_label:
        run = paragraph.add_run(bold_label)
        run.bold = True
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
        paragraph.add_run(text)
    else:
        paragraph.add_run(text)
    for run in paragraph.runs:
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    return paragraph


def add_small_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], title: str) -> None:
    add_heading(doc, title, level=2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=8.5)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            set_cell_text(cells[i], value, size=8.0)
    set_table_widths(table, widths)
    doc.add_paragraph()


def font_for_chart(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size, index=0)
    return ImageFont.load_default()


def draw_category_chart(records: list[Record]) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / "category_counts.png"
    passed = [record for record in records if record.status == "passed"]
    counts = Counter(record.category for record in passed)
    ordered = sorted(counts.items(), key=lambda item: CATEGORY_ORDER.get(item[0], 99))
    width, height = 1900, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font_for_chart(42)
    label_font = font_for_chart(24)
    small_font = font_for_chart(22)
    draw.text((60, 40), "Google Codelabs PDF 主分類分布", fill="#202124", font=title_font)
    left, top, bar_h, gap = 760, 120, 42, 24
    max_count = max(counts.values()) if counts else 1
    for i, (category, count) in enumerate(ordered):
        y = top + i * (bar_h + gap)
        label = category
        if len(label) > 38:
            label = label.replace(" / ", " /\n", 1)
        draw.multiline_text((60, y + 2), label, fill="#202124", font=label_font, spacing=2)
        bar_w = int((width - left - 180) * count / max_count)
        draw.rounded_rectangle((left, y, left + bar_w, y + bar_h), radius=8, fill="#1A73E8")
        draw.text((left + bar_w + 16, y + 7), str(count), fill="#202124", font=small_font)
    draw.text((60, height - 70), "分類為依標題、URL slug、描述文字推論；每列仍保留原始 PDF 檔名與來源 URL。", fill="#5F6368", font=small_font)
    image.save(path)
    return path


def draw_ai_timeline_chart(records: list[Record]) -> Path:
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    path = CHART_DIR / "ai_timeline.png"
    ai_records = [record for record in records if record.category == "AI / Gemini / Agent / Vertex AI" and record.updated_date]
    month_counts: dict[str, Counter[str]] = defaultdict(Counter)
    for record in ai_records:
        month_counts[record.updated_date[:7]][record.ai_subtopic or "其他 AI / GenAI"] += 1
    months = sorted(month_counts)
    topics = [
        "Gemini App 開發",
        "Vertex AI / Gemini API",
        "RAG / grounding / 搜尋",
        "BigQuery AI / 資料代理",
        "ADK / 多代理 / Agentic Apps",
        "Enterprise Agent Platform / Gateway",
        "AI 評估 / 觀測",
        "AI 安全 / 治理",
        "MCP / Workspace 代理",
        "其他 AI / GenAI",
    ]
    palette = ["#1A73E8", "#34A853", "#FBBC04", "#A142F4", "#F4511E", "#00ACC1", "#5F6368", "#C5221F", "#188038", "#9AA0A6"]
    width, height = 1700, 900
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font_for_chart(42)
    label_font = font_for_chart(26)
    small_font = font_for_chart(20)
    draw.text((60, 40), "AI / Gemini / Agent 類 Codelabs：Updated 時間演進", fill="#202124", font=title_font)
    x0, y0, chart_w, chart_h = 120, 160, 1050, 560
    max_total = max((sum(month_counts[m].values()) for m in months), default=1)
    bar_w = 190
    for i, month in enumerate(months):
        x = x0 + i * 300
        y_bottom = y0 + chart_h
        total = sum(month_counts[month].values())
        current_bottom = y_bottom
        for topic, color in zip(topics, palette):
            count = month_counts[month][topic]
            if not count:
                continue
            h = int(chart_h * count / max_total)
            draw.rectangle((x, current_bottom - h, x + bar_w, current_bottom), fill=color)
            current_bottom -= h
        draw.text((x + 25, y_bottom + 20), month, fill="#202124", font=label_font)
        draw.text((x + 58, current_bottom - 34), str(total), fill="#202124", font=label_font)
    legend_x, legend_y = 1220, 150
    for i, (topic, color) in enumerate(zip(topics, palette)):
        y = legend_y + i * 58
        draw.rectangle((legend_x, y, legend_x + 28, y + 28), fill=color)
        draw.text((legend_x + 42, y - 2), topic, fill="#202124", font=small_font)
    draw.text((60, height - 70), "注意：時間為卡片 Updated 日期，用來觀察這批學習材料的更新重點，不等同產品首次發布日。", fill="#5F6368", font=small_font)
    image.save(path)
    return path


def export_csv(records: list[Record]) -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    fields = [
        "index",
        "slug",
        "category",
        "level",
        "updated_date",
        "duration_text",
        "title",
        "pdf_filename",
        "status",
        "step_count",
        "pdf_page_count",
        "url",
    ]
    sorted_records = sort_records(records)
    with OUT_CSV.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        for record in sorted_records:
            writer.writerow(
                {
                    "index": record.index,
                    "slug": record.slug,
                    "category": record.category,
                    "level": record.level,
                    "updated_date": record.updated_date,
                    "duration_text": record.duration_text,
                    "title": record.title or record.slug,
                    "pdf_filename": record.pdf_filename,
                    "status": record.status,
                    "step_count": record.step_count,
                    "pdf_page_count": record.pdf_page_count,
                    "url": record.url,
                }
            )


def sort_records(records: list[Record]) -> list[Record]:
    return sorted(
        records,
        key=lambda record: (
            CATEGORY_ORDER.get(record.category, 99),
            LEVEL_ORDER.get(record.level, 9),
            record.updated_date or "0000-00-00",
            record.index,
        ),
    )


def build_doc(records: list[Record]) -> None:
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    category_chart = draw_category_chart(records)
    timeline_chart = draw_ai_timeline_chart(records)
    export_csv(records)

    passed = [record for record in records if record.status == "passed"]
    failed = [record for record in records if record.status != "passed"]
    category_counts = Counter(record.category for record in passed)
    level_counts = Counter(record.level for record in passed)
    ai_records = [record for record in passed if record.category == "AI / Gemini / Agent / Vertex AI"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.style = styles["Title"]
    title_run = title.add_run("Google Codelabs PDF 學習地圖與完整索引")
    title_run.font.name = "Calibri"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_body_paragraph(
        doc,
        f"資料範圍：Google Codelabs 首頁 zh-tw 清單共 {len(records)} 筆；本地 PDF 已通過驗證 {len(passed)} 份，未產生 PDF {len(failed)} 筆。",
        bold_label="摘要：",
    )
    add_body_paragraph(
        doc,
        "分類依據是每個 codelab 的標題、URL slug、首頁描述與 PDF 驗證報表；Updated 日期取自 Google Codelabs 首頁卡片，不等同原始首發日。",
        bold_label="資料口徑：",
    )
    add_body_paragraph(
        doc,
        "閱讀順序建議先看同分類內的初階，再看中階、進階、高階；AI 類建議優先走 Gemini 基礎、Vertex AI / RAG、ADK / 多代理、評估與安全、Enterprise Agent Platform。",
        bold_label="怎麼看：",
    )

    summary_rows = [
        ["PDF 通過驗證", str(len(passed))],
        ["未產生 PDF", str(len(failed))],
        ["AI / Gemini / Agent 類", str(len(ai_records))],
        ["初階 / 中階 / 進階 / 高階", f"{level_counts['初階']} / {level_counts['中階']} / {level_counts['進階']} / {level_counts['高階']}"],
    ]
    add_small_table(doc, ["指標", "數量"], summary_rows, [2.3, 4.2], "快速總覽")

    doc.add_picture(str(category_chart), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    category_rows = []
    for category, _order in sorted(CATEGORY_ORDER.items(), key=lambda item: item[1]):
        rows_for_category = [record for record in passed if record.category == category]
        if not rows_for_category:
            continue
        counts = Counter(record.level for record in rows_for_category)
        category_rows.append(
            [
                category,
                str(len(rows_for_category)),
                str(counts["初階"]),
                str(counts["中階"]),
                str(counts["進階"]),
                str(counts["高階"]),
            ]
        )
    add_small_table(doc, ["主分類", "總數", "初階", "中階", "進階", "高階"], category_rows, [2.2, 0.65, 0.65, 0.65, 0.65, 0.65], "分類數量總表")

    add_heading(doc, "表 1：完整 PDF 分類索引", level=1)
    add_body_paragraph(
        doc,
        "這張表列出每一份 PDF 的檔名、分類、學習階段與 Updated 日期。排序方式是主分類 → 初階到高階 → Updated 日期 → 原始序號。",
    )
    headers = ["序", "主分類", "階段", "Updated", "時長", "標題", "PDF 檔名"]
    widths = [0.35, 0.9, 0.43, 0.68, 0.45, 2.2, 1.49]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=7.4)
    set_repeat_table_header(table.rows[0])

    for record in sort_records(records):
        cells = table.add_row().cells
        status_note = "" if record.status == "passed" else "（未產 PDF）"
        values = [
            f"{record.index:03d}",
            record.category.replace(" / ", "\n"),
            record.level,
            record.updated_date,
            record.duration_text,
            record.title or record.slug,
            record.pdf_filename or status_note,
        ]
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=6.7 if i in (1, 5, 6) else 7.0, color="9B1C1C" if record.status != "passed" else "202124")
        if record.status != "passed":
            for cell in cells:
                set_cell_shading(cell, "FCE8E6")
    set_table_widths(table, widths)

    doc.add_page_break()
    add_heading(doc, "表 2：Google AI 主題時間演進", level=1)
    add_body_paragraph(
        doc,
        "這張表只看 AI / Gemini / Agent / Vertex AI 類 codelab，依 Google Codelabs 卡片的 Updated 日期分月觀察重點。因所有卡片目前集中在 2026-03 到 2026-05，這裡呈現的是這批教材的更新節奏，不是 Google AI 產品完整歷史。",
    )

    month_rows = []
    ai_by_month: dict[str, list[Record]] = defaultdict(list)
    for record in ai_records:
        if record.updated_date:
            ai_by_month[record.updated_date[:7]].append(record)
    for month in sorted(ai_by_month):
        rows_for_month = ai_by_month[month]
        topic_counts = Counter(record.ai_subtopic for record in rows_for_month)
        category_focus = "、".join(topic for topic, _ in topic_counts.most_common(4))
        reps = "；".join(record.title for record in rows_for_month[:3])
        if month == "2026-03":
            interpretation = "教材刷新集中在 Gemini / Vertex AI / Cloud Run / BigQuery，適合建立生成式 AI 應用與資料基礎。"
        elif month == "2026-04":
            interpretation = "重心轉向 production-ready AI：ADK、多代理、評估、安全、觀測與治理。"
        elif month == "2026-05":
            interpretation = "最新材料偏向 Enterprise Agent Platform、Agent Gateway、MCP、Workspace 代理與事件驅動代理。"
        else:
            interpretation = "以當月新增或更新教材觀察主要技術焦點。"
        month_rows.append([month, str(len(rows_for_month)), category_focus, reps, interpretation])

    add_small_table(
        doc,
        ["時間區間", "AI 案子數", "主要焦點", "代表 codelab", "給你的學習判讀"],
        month_rows,
        [0.75, 0.55, 1.45, 1.9, 1.85],
        "表 2：AI 教材更新時間軸",
    )
    doc.add_picture(str(timeline_chart), width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "建議學習路線", level=1)
    path_rows = [
        ["1", "初階", "Gemini 基礎、開始使用 API、簡單 Cloud Run / Firebase / Flutter 範例", "先建立模型呼叫、資料輸入、部署入口的共同語言。"],
        ["2", "中階", "Vertex AI、BigQuery AI、RAG、生成式 UI、資料代理入門", "開始把模型接到資料、搜尋、前端與後端工作流。"],
        ["3", "進階", "ADK、Agentic Apps、多代理、MCP、Workspace 代理", "開始學代理如何拆任務、調工具、連企業資料與外部服務。"],
        ["4", "高階", "評估、觀測、安全、Model Armor、Agent Gateway、Enterprise Agent Platform", "最後補上 production-grade 的治理、安全、網路與可觀測性。"],
    ]
    add_small_table(doc, ["順序", "階段", "先學什麼", "為什麼"], path_rows, [0.45, 0.7, 2.45, 2.9], "從初階到高階的閱讀順序")

    if failed:
        add_heading(doc, "未產生 PDF 的項目", level=1)
        failed_rows = [[f"{record.index:03d}", record.slug, record.url, "頁面沒有偵測到 codelab step，避免產生假 PDF"] for record in failed]
        add_small_table(doc, ["序", "slug", "URL", "原因"], failed_rows, [0.4, 1.5, 2.7, 1.9], "失敗項目")

    doc.save(OUT_DOCX)


def main() -> None:
    records = load_records()
    build_doc(records)
    print(f"docx={OUT_DOCX}")
    print(f"csv={OUT_CSV}")


if __name__ == "__main__":
    main()
